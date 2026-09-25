"""Attach revenue and employee counts to the CIW Wales domiciliary screen.

Reads the CIW workbook, looks up each row's Companies House number in the
local database (company_enrichment), and writes a copy of the workbook with
financial columns appended to every sheet that has a "Company Number" column.

Enrich first if needed:
    python -m ch_bulk enrich-financials --mode list --ids <comma-separated>

Run: python scripts/export_ciw_wales_financials.py
"""

from __future__ import annotations

import argparse
import sys
from copy import copy
from pathlib import Path

import duckdb
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

from export_complex_care_merged import (  # noqa: E402
    ESTIMATED_REVENUE_SOURCES,
    norm_company_number,
    usable_employee_count,
)

from ch_bulk.companies_house.revenue_model import estimate, load_bands  # noqa: E402
from ch_bulk.core.paths import DEFAULT_DATA_DIR, DEFAULT_DB_PATH, revenue_bands_path  # noqa: E402

SOURCE_XLSX = Path("/Users/nitinkum/Downloads/CIW_Wales_Domiciliary_Screen_Classified.xlsx")

NEW_COLUMNS = [
    "CH Company Status",
    "Employee Count",
    "Revenue (filed accounts)",
    "Revenue Source",
    "Revenue (estimated from employees)",
    "Accounts Period End",
    "Accounts Category",
    "Financials Note",
]

FACTS_SQL = """
SELECT
    c.company_number,
    c.company_status,
    c.accounts_category,
    e.company_number IS NOT NULL AS enriched,
    e.employee_count,
    e.revenue,
    e.revenue_source,
    e.filing_period_end
FROM companies c
LEFT JOIN company_enrichment e USING (company_number)
WHERE c.company_number IN (SELECT chn FROM wanted)
"""


def load_facts(con, numbers: list[str]) -> dict[str, dict]:
    con.execute("CREATE OR REPLACE TEMP TABLE wanted (chn VARCHAR)")
    con.executemany("INSERT INTO wanted VALUES (?)", [(n,) for n in numbers])
    cur = con.execute(FACTS_SQL)
    cols = [d[0] for d in cur.description]
    return {row[0]: dict(zip(cols, row)) for row in cur.fetchall()}


def financial_values(chn: str | None, facts: dict[str, dict], bands) -> list:
    if chn is None:
        return [None] * (len(NEW_COLUMNS) - 1) + ["No Companies House number"]
    f = facts.get(chn)
    if f is None:
        return [None] * (len(NEW_COLUMNS) - 1) + ["Not found in local CH database"]

    source = f["revenue_source"]
    filed_revenue = f["revenue"] if source not in ESTIMATED_REVENUE_SOURCES else None
    employees, rejected = usable_employee_count(f["employee_count"])

    if not f["enriched"]:
        note = "Not enriched"
    elif filed_revenue is None and employees is None:
        note = "No revenue or employees in latest filed accounts"
    elif filed_revenue is None:
        note = "Revenue not disclosed (small/micro accounts)"
    else:
        note = None
    if rejected is not None:
        note = f"Implausible employee figure dropped ({rejected})"

    return [
        f["company_status"],
        employees,
        filed_revenue,
        source,
        estimate(employees, bands),
        f["filing_period_end"],
        f["accounts_category"],
        note,
    ]


def annotate_sheet(ws, facts: dict[str, dict], bands) -> int:
    headers = [c.value for c in ws[1]]
    chn_idx = headers.index("Company Number")
    start_col = ws.max_column + 1
    header_cell = ws.cell(row=1, column=1)

    for offset, name in enumerate(NEW_COLUMNS):
        cell = ws.cell(row=1, column=start_col + offset, value=name)
        cell.font = copy(header_cell.font)
        cell.fill = copy(header_cell.fill)
        cell.alignment = copy(header_cell.alignment)
        cell.border = copy(header_cell.border)
        ws.column_dimensions[cell.column_letter].width = max(14, len(name) + 2)

    matched = 0
    for row in range(2, ws.max_row + 1):
        chn = norm_company_number(ws.cell(row=row, column=chn_idx + 1).value)
        values = financial_values(chn, facts, bands)
        if values[1] is not None or values[2] is not None:
            matched += 1
        for offset, value in enumerate(values):
            cell = ws.cell(row=row, column=start_col + offset, value=value)
            if NEW_COLUMNS[offset].startswith("Revenue (") and value is not None:
                cell.number_format = "£#,##0"
            elif NEW_COLUMNS[offset] == "Accounts Period End" and value is not None:
                cell.number_format = "yyyy-mm-dd"
    return matched


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--source", type=Path, default=SOURCE_XLSX)
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--db-path", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    args = parser.parse_args()
    out = args.out or args.source.with_name(f"{args.source.stem}_with_financials.xlsx")

    wb = openpyxl.load_workbook(args.source)
    sheets = [ws for ws in wb if "Company Number" in [c.value for c in ws[1]]]
    numbers = sorted(
        {
            n
            for ws in sheets
            for n in (
                norm_company_number(v)
                for (v,) in ws.iter_rows(
                    min_row=2,
                    min_col=[c.value for c in ws[1]].index("Company Number") + 1,
                    max_col=[c.value for c in ws[1]].index("Company Number") + 1,
                    values_only=True,
                )
            )
            if n
        }
    )

    con = duckdb.connect(str(args.db_path), read_only=True)
    facts = load_facts(con, numbers)
    con.close()
    bands = load_bands(revenue_bands_path(args.data_dir))

    for ws in sheets:
        matched = annotate_sheet(ws, facts, bands)
        print(f"{ws.title}: {ws.max_row - 1} rows, {matched} with revenue or employees")

    wb.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
