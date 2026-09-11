"""CQC Homecare Gap Analysis — Task 1 + Task 2.

Reads the input workbook, filters/matches/reclassifies, writes a new Excel.

Run: .venv/bin/python scripts/cqc_homecare_gap.py
"""

from __future__ import annotations

import argparse
import re
from datetime import date
from pathlib import Path

import duckdb
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz

INPUT_XLSX = Path(
    "/Users/nitinkum/Downloads/13062026_Potentially_missing_homecare_-_CQC_data.xlsx"
)
DB_PATH = Path("data/db/ch_bulk.duckdb")

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TIER1_FILL = PatternFill("solid", fgColor="C6EFCE")
TIER2_FILL = PatternFill("solid", fgColor="FFEB9C")
TIER3_FILL = PatternFill("solid", fgColor="F4CCCC")
EXCLUDED_FILL = PatternFill("solid", fgColor="D9D9D9")
INCLUDE_FILL = PatternFill("solid", fgColor="C6EFCE")
EXCLUDE_FILL = PatternFill("solid", fgColor="FFC7CE")
REVIEW_FILL = PatternFill("solid", fgColor="FFEB9C")

NOISE_TOKENS = {"ltd", "limited", "plc", "llp", "the", "uk"}
FUZZY_THRESHOLD = 85

EXCLUDED_OWNERSHIP_TYPES = {"Individual", "Partnership", "NHS Body"}

NON_INDEPENDENT_NAME_SIGNALS = [
    "charity", "trust", "council", "borough council", "county council",
    "nhs", "foundation", "association", "society",
]

EXCLUDED_COMPANY_TYPES = {
    "Registered Society",
    "Community Interest Company",
    "Charitable Incorporated Organisation",
    "Scottish Charitable Incorporated Organisation",
    "Royal Charter Company",
    "Industrial and Provident Society",
    "Other company type",
    "Other Company Type",
    "Overseas Entity",
}
EXCLUDED_COMPANY_TYPE_PREFIXES = [
    "PRI/LTD BY GUAR/NSC",
    "PRI/LBG/NSC",
]

AUTO_PASS_COMPANY_TYPES = {
    "Private Limited Company",
    "Limited Liability Partnership",
}

MANUAL_REVIEW_COMPANY_TYPES = {
    "Private Unlimited Company",
    "Limited Partnership",
    "Public Limited Company",
}


# --------------- normalisation helpers ---------------


def normalize_name(name: str) -> str:
    cleaned = re.sub(r"[^a-z0-9 ]+", " ", (name or "").lower())
    tokens = [t for t in cleaned.split() if t and t not in NOISE_TOKENS]
    return " ".join(tokens).strip()


def normalize_postcode(pc: str | None) -> str | None:
    if not pc:
        return None
    cleaned = re.sub(r"\s+", "", str(pc).upper())
    return cleaned or None


def normalize_company_number(val: object) -> str | None:
    text = re.sub(r"\s+", "", str(val or "").strip().upper())
    if not text:
        return None
    if text.isdigit():
        return text.zfill(8)
    return text


def postcode_outward(pc: str | None) -> str | None:
    normed = normalize_postcode(pc)
    if not normed:
        return None
    if len(normed) <= 4:
        return normed
    return normed[:-3]


def has_name_signal(name: str) -> str | None:
    lower = (name or "").lower()
    for signal in NON_INDEPENDENT_NAME_SIGNALS:
        if signal in lower:
            return signal
    return None


def classify_company_type(ct: str | None) -> tuple[str, str]:
    if not ct:
        return "review", "no company_type available"
    if ct in AUTO_PASS_COMPANY_TYPES:
        return "include", f"company_type={ct}"
    if ct in EXCLUDED_COMPANY_TYPES:
        return "exclude", f"company_type={ct}"
    for prefix in EXCLUDED_COMPANY_TYPE_PREFIXES:
        if ct.startswith(prefix):
            return "exclude", f"company_type starts with {prefix}"
    if ct in MANUAL_REVIEW_COMPANY_TYPES:
        return "review", f"company_type={ct} (rare commercial form)"
    return "review", f"unknown company_type={ct}"


# --------------- Excel loaders ---------------


def load_potentially_missing(wb) -> list[dict]:
    ws = wb["Potentially missing homecare"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        provider_name = str(row[0] or "").strip()
        if not provider_name:
            continue
        rows.append({
            "provider_name": provider_name,
            "provider_id": str(row[1] or "").strip(),
            "num_locations": row[2],
            "primary_location": str(row[3] or "").strip(),
            "address": str(row[4] or "").strip(),
            "postcode": str(row[5] or "").strip(),
            "phone": row[6],
            "website": str(row[7] or "").strip(),
            "service_types": str(row[8] or "").strip(),
            "local_authority": str(row[11] or "").strip(),
            "region": str(row[12] or "").strip(),
        })
    return rows


def load_database_tab(wb) -> list[dict]:
    ws = wb["database"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        company_number_raw = row[2]
        if company_number_raw is None:
            continue
        rows.append({
            "tier": row[0],
            "total_score": row[1],
            "company_number_raw": company_number_raw,
            "company_number": normalize_company_number(company_number_raw),
            "company_name": str(row[3] or "").strip(),
            "company_name_normalized": normalize_name(str(row[3] or "")),
            "status": row[4],
            "town": str(row[5] or "").strip(),
            "postcode": str(row[6] or "").strip(),
            "postcode_normalized": normalize_postcode(str(row[6] or "")),
            "region": str(row[7] or "").strip(),
            "classification": str(row[8] or "").strip(),
            "revenue": row[10],
            "employees": row[13],
            "avg_director_age": row[14],
            "oldest_director": row[15],
            "directors_60_plus": row[16],
            "all_directors_60_plus": row[17],
            "cqc_provider_name": str(row[18] or "").strip(),
            "cqc_rating": row[19],
            "cqc_domiciliary": row[22],
            "cqc_supported_living": row[23],
            "cqc_nursing": row[24],
            "cqc_no_nursing": row[25],
            "excluded_by_workbook": row[28],
        })
    return rows


# --------------- DB helpers ---------------


def build_cqc_enrichment_lookup(con, provider_ids: list[str]) -> dict[str, dict]:
    if not provider_ids:
        return {}
    placeholders = ", ".join(["?"] * len(provider_ids))
    rows = con.execute(f"""
        SELECT
            pe.provider_id,
            pe.ownership_type,
            pe.brand_name,
            pe.charity_number,
            pe.companies_house_number,
            pe.company_name
        FROM cqc_providers_enriched pe
        WHERE pe.provider_id IN ({placeholders})
    """, provider_ids).fetchall()
    return {
        str(r[0]): {
            "ownership_type": r[1],
            "brand_name": r[2],
            "charity_number": r[3],
            "companies_house_number": normalize_company_number(r[4]),
            "cqc_company_name": r[5],
        }
        for r in rows
    }


def build_hsca_ch_lookup(con, provider_ids: list[str]) -> dict[str, str]:
    if not provider_ids:
        return {}
    placeholders = ", ".join(["?"] * len(provider_ids))
    rows = con.execute(f"""
        SELECT DISTINCT provider_id, provider_companies_house_number
        FROM cqc_hsca_locations
        WHERE provider_id IN ({placeholders})
          AND provider_companies_house_number IS NOT NULL
          AND TRIM(provider_companies_house_number) != ''
    """, provider_ids).fetchall()
    result: dict[str, str] = {}
    for pid, ch_num in rows:
        normed = normalize_company_number(ch_num)
        if normed and str(pid) not in result:
            result[str(pid)] = normed
    return result


def build_company_lookup(con, company_numbers: list[str]) -> dict[str, dict]:
    if not company_numbers:
        return {}
    placeholders = ", ".join(["?"] * len(company_numbers))
    rows = con.execute(f"""
        SELECT company_number, company_name, company_type, postcode, company_status
        FROM companies
        WHERE company_number IN ({placeholders})
    """, company_numbers).fetchall()
    return {
        str(r[0]): {
            "company_name": r[1],
            "company_type": r[2],
            "postcode": r[3],
            "company_status": r[4],
        }
        for r in rows
    }


def fuzzy_search_company(con, name: str, postcode: str | None) -> dict | None:
    outward = postcode_outward(postcode)
    if not outward or not name:
        return None
    normed_name = normalize_name(name)
    if not normed_name:
        return None
    rows = con.execute("""
        SELECT company_number, company_name, company_type, postcode
        FROM companies
        WHERE company_status = 'Active'
          AND is_active = TRUE
          AND CASE
              WHEN LENGTH(REGEXP_REPLACE(UPPER(COALESCE(postcode, '')), '\\s+', '', 'g')) <= 4
                  THEN REGEXP_REPLACE(UPPER(COALESCE(postcode, '')), '\\s+', '', 'g')
              ELSE SUBSTR(REGEXP_REPLACE(UPPER(COALESCE(postcode, '')), '\\s+', '', 'g'), 1,
                   LENGTH(REGEXP_REPLACE(UPPER(COALESCE(postcode, '')), '\\s+', '', 'g')) - 3)
          END = ?
    """, [outward]).fetchall()

    best = None
    best_ratio = 0
    for r in rows:
        cn = normalize_name(r[1])
        ratio = int(round(fuzz.token_set_ratio(normed_name, cn)))
        if ratio > best_ratio:
            best_ratio = ratio
            best = {
                "company_number": r[0],
                "company_name": r[1],
                "company_type": r[2],
                "postcode": r[3],
                "ratio": ratio,
            }
    if best and best_ratio >= FUZZY_THRESHOLD:
        return best
    return None


def build_nursing_check(con, company_numbers: list[str]) -> dict[str, bool]:
    if not company_numbers:
        return {}
    placeholders = ", ".join(["?"] * len(company_numbers))
    rows = con.execute(f"""
        SELECT DISTINCT m.company_number
        FROM ch_cqc_matches m
        JOIN cqc_locations cl ON cl.provider_id = m.cqc_provider_id
        WHERE m.company_number IN ({placeholders})
          AND cl.service_types LIKE '%Community services - Nursing%'
    """, company_numbers).fetchall()
    return {str(r[0]): True for r in rows}


def build_enrichment_fallback(con, company_numbers: list[str]) -> dict[str, dict]:
    if not company_numbers:
        return {}
    placeholders = ", ".join(["?"] * len(company_numbers))
    rows = con.execute(f"""
        SELECT company_number, max_director_age, directors_over_60, all_directors_60_plus
        FROM company_enrichment
        WHERE company_number IN ({placeholders})
    """, company_numbers).fetchall()
    return {
        str(r[0]): {
            "max_director_age": r[1],
            "directors_over_60": r[2],
            "all_directors_60_plus": r[3],
        }
        for r in rows
    }


# --------------- Task 1 ---------------


def run_task1(missing_rows: list[dict], db_rows: list[dict], con) -> tuple[list[dict], list[dict]]:
    provider_ids = [r["provider_id"] for r in missing_rows if r["provider_id"]]

    cqc_enrichment = build_cqc_enrichment_lookup(con, provider_ids)
    hsca_ch = build_hsca_ch_lookup(con, provider_ids)

    # Build database lookup indices
    db_by_number: dict[str, dict] = {}
    db_by_name_pc: dict[tuple[str, str | None], dict] = {}
    for dbr in db_rows:
        if dbr["company_number"]:
            db_by_number[dbr["company_number"]] = dbr
        key = (dbr["company_name_normalized"], dbr["postcode_normalized"])
        if key[0]:
            db_by_name_pc[key] = dbr

    # Collect all CH numbers we find, then batch-lookup companies
    ch_numbers_to_lookup: set[str] = set()
    for row in missing_rows:
        pid = row["provider_id"]
        enriched = cqc_enrichment.get(pid, {})
        ch_num = enriched.get("companies_house_number") or hsca_ch.get(pid)
        if ch_num:
            ch_numbers_to_lookup.add(ch_num)

    companies_data = build_company_lookup(con, list(ch_numbers_to_lookup))

    # Process each row
    results = []
    review_queue = []

    for row in missing_rows:
        pid = row["provider_id"]
        enriched = cqc_enrichment.get(pid, {})
        ownership_type = enriched.get("ownership_type")
        brand_name = enriched.get("brand_name")
        charity_number = enriched.get("charity_number")

        result = {
            "provider_id": pid,
            "provider_name": row["provider_name"],
            "postcode": row["postcode"],
            "service_types": row["service_types"],
            "local_authority": row["local_authority"],
            "region": row["region"],
            "ownership_type": ownership_type or "",
            "brand_name": brand_name or "",
            "charity_number": charity_number or "",
            "candidate_company_number": "",
            "candidate_company_name": "",
            "company_type": "",
            "match_source": "",
            "independence_decision": "",
            "decision_reason": "",
            "already_in_database": "",
            "database_row_tier": "",
        }

        # --- CQC-level exclusions (before CH lookup) ---

        early_exclude = False
        if ownership_type in EXCLUDED_OWNERSHIP_TYPES:
            result["independence_decision"] = "exclude"
            result["decision_reason"] = f"ownership_type={ownership_type}"
            early_exclude = True

        elif charity_number:
            result["independence_decision"] = "exclude"
            result["decision_reason"] = f"has charity_number={charity_number}"
            early_exclude = True

        else:
            name_signal = has_name_signal(row["provider_name"])
            if name_signal:
                result["independence_decision"] = "exclude"
                result["decision_reason"] = f"name contains '{name_signal}'"
                early_exclude = True

            elif brand_name:
                result["independence_decision"] = "exclude"
                result["decision_reason"] = f"franchise/chain brand={brand_name}"
                early_exclude = True

        # --- Resolve CH number (always, even for excluded rows) ---

        ch_num = enriched.get("companies_house_number")
        match_source = "cqc_api" if ch_num else None

        if not ch_num:
            ch_num = hsca_ch.get(pid)
            if ch_num:
                match_source = "hsca"

        company_info = None
        if ch_num:
            company_info = companies_data.get(ch_num)

        if not early_exclude:
            # Local exact match fallback
            if not ch_num:
                normed_name = normalize_name(row["provider_name"])
                normed_pc = normalize_postcode(row["postcode"])
                exact_rows = con.execute("""
                    SELECT company_number, company_name, company_type, postcode
                    FROM companies
                    WHERE company_status = 'Active' AND is_active = TRUE
                      AND LOWER(REGEXP_REPLACE(company_name, '[^a-zA-Z0-9 ]', ' ', 'g')) = ?
                      AND UPPER(REGEXP_REPLACE(postcode, '\\s+', '', 'g')) = ?
                    LIMIT 1
                """, [normed_name, normed_pc or ""]).fetchall()
                if exact_rows:
                    ch_num = exact_rows[0][0]
                    match_source = "local_exact"
                    company_info = {
                        "company_name": exact_rows[0][1],
                        "company_type": exact_rows[0][2],
                        "postcode": exact_rows[0][3],
                    }

            # Fuzzy fallback
            if not ch_num:
                fuzzy_hit = fuzzy_search_company(con, row["provider_name"], row["postcode"])
                if fuzzy_hit:
                    ch_num = fuzzy_hit["company_number"]
                    match_source = "local_fuzzy"
                    company_info = fuzzy_hit

        if ch_num:
            result["candidate_company_number"] = ch_num
            result["candidate_company_name"] = (company_info or {}).get("company_name", "")
            result["company_type"] = (company_info or {}).get("company_type", "")
            result["match_source"] = match_source or ""

        # --- Company-type screening (only for non-early-excluded rows) ---

        if not early_exclude:
            if company_info:
                ct_decision, ct_reason = classify_company_type(company_info.get("company_type"))
                result["independence_decision"] = ct_decision
                result["decision_reason"] = ct_reason
            elif not ch_num:
                result["independence_decision"] = "review"
                result["decision_reason"] = "no CH number found"
                result["match_source"] = "unresolved"
            else:
                result["independence_decision"] = "review"
                result["decision_reason"] = "CH number found but company not in local DB"

        # --- Cross-check against database tab (always) ---

        db_match = None
        if ch_num and ch_num in db_by_number:
            db_match = db_by_number[ch_num]
        if not db_match:
            key = (normalize_name(row["provider_name"]), normalize_postcode(row["postcode"]))
            if key[0] and key in db_by_name_pc:
                db_match = db_by_name_pc[key]

        if db_match:
            result["already_in_database"] = "Yes"
            result["database_row_tier"] = db_match.get("tier", "")
        else:
            result["already_in_database"] = "No"

        # Route fuzzy matches to review queue
        if not early_exclude and match_source == "local_fuzzy":
            ratio = (company_info or {}).get("ratio", 0)
            result["decision_reason"] += f" (fuzzy ratio={ratio})"
            review_queue.append(result)
        else:
            results.append(result)

    return results, review_queue


# --------------- Task 2 ---------------


def run_task2(db_rows: list[dict], con) -> list[dict]:
    target_rows = [
        r for r in db_rows
        if r["tier"] == "Excluded"
        and r["classification"] in ("Domiciliary Care", "Domiciliary Care + Supported Living")
    ]

    company_numbers = [r["company_number"] for r in target_rows if r["company_number"]]
    enrichment_fallback = build_enrichment_fallback(con, company_numbers)
    nursing_flags = build_nursing_check(con, company_numbers)

    results = []
    for row in target_rows:
        cn = row["company_number"]
        classification = row["classification"]

        # Determine director age
        oldest = row["oldest_director"]
        dirs_60 = row["directors_60_plus"]
        data_quality = ""

        if oldest is None and cn:
            fb = enrichment_fallback.get(cn, {})
            oldest = fb.get("max_director_age")
            dirs_60 = fb.get("directors_over_60")
            if oldest is not None:
                data_quality = "age from company_enrichment (not in workbook)"
            else:
                data_quality = "no director age data available"

        # Apply reclassification rules
        if classification == "Domiciliary Care + Supported Living":
            new_tier = "Tier 3"
            tier_reason = "Domiciliary Care + Supported Living → Tier 3"
        elif classification == "Domiciliary Care":
            has_60_plus = False
            if oldest is not None:
                try:
                    has_60_plus = float(oldest) >= 60
                except (ValueError, TypeError):
                    pass
            if not has_60_plus and dirs_60 is not None:
                try:
                    has_60_plus = int(dirs_60) > 0
                except (ValueError, TypeError):
                    pass

            if has_60_plus:
                new_tier = "Tier 1"
                tier_reason = f"Domiciliary Care, oldest director={oldest} (≥60)"
            else:
                new_tier = "Tier 2"
                tier_reason = f"Domiciliary Care, oldest director={oldest} (<60 or unknown)"
        else:
            new_tier = "Excluded"
            tier_reason = "unexpected classification"

        has_nursing = nursing_flags.get(cn, False) if cn else False

        results.append({
            "company_number": cn or str(row["company_number_raw"]),
            "company_name": row["company_name"],
            "postcode": row["postcode"],
            "classification": classification,
            "original_tier": "Excluded",
            "oldest_director": oldest,
            "directors_60_plus": dirs_60,
            "new_tier": new_tier,
            "has_non_residential_nursing": "Yes" if has_nursing else "No",
            "tier_reason": tier_reason,
            "data_quality_note": data_quality,
        })

    return results


# --------------- Excel writer ---------------


def write_table(ws, headers: list[str], rows: list[dict], key_order: list[str]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "A2"

    for row_dict in rows:
        ws.append([row_dict.get(k, "") for k in key_order])

    for col_i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_i)].width = max(
            12, min(45, len(header) + 4)
        )
    ws.auto_filter.ref = ws.dimensions


def style_decision_column(ws, col_header: str):
    headers = [c.value for c in ws[1]]
    if col_header not in headers:
        return
    idx = headers.index(col_header) + 1
    letter = get_column_letter(idx)
    for cell in ws[letter][1:]:
        v = str(cell.value or "").lower()
        if v == "include":
            cell.fill = INCLUDE_FILL
        elif v == "exclude":
            cell.fill = EXCLUDE_FILL
        elif v == "review":
            cell.fill = REVIEW_FILL


def style_tier_column(ws, col_header: str = "new_tier"):
    headers = [c.value for c in ws[1]]
    if col_header not in headers:
        return
    idx = headers.index(col_header) + 1
    letter = get_column_letter(idx)
    for cell in ws[letter][1:]:
        v = str(cell.value or "")
        if v == "Tier 1":
            cell.fill = TIER1_FILL
        elif v == "Tier 2":
            cell.fill = TIER2_FILL
        elif v == "Tier 3":
            cell.fill = TIER3_FILL
        elif v == "Excluded":
            cell.fill = EXCLUDED_FILL


def build_summary(ws, task1_results, task1_review, task2_results):
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 15

    title_font = Font(bold=True, size=14, color="1F4E79")
    section_font = Font(bold=True, size=12, color="1F4E79")

    ws.append([f"CQC Homecare Gap Analysis ({date.today().isoformat()})"])
    ws.cell(1, 1).font = title_font
    ws.append([])

    ws.append(["Task 1 — Potentially missing homecare providers"])
    ws.cell(ws.max_row, 1).font = section_font

    all_t1 = task1_results + task1_review
    ws.append(["Total populated input rows", len(all_t1)])
    include = sum(1 for r in all_t1 if r["independence_decision"] == "include")
    exclude = sum(1 for r in all_t1 if r["independence_decision"] == "exclude")
    review = sum(1 for r in all_t1 if r["independence_decision"] == "review")
    ws.append(["  Include (independent private company)", include])
    ws.append(["  Exclude (non-independent)", exclude])
    ws.append(["  Review needed", review])
    ws.append(["  — of which fuzzy matches (in review tab)", len(task1_review)])

    already_in = sum(1 for r in all_t1 if r["already_in_database"] == "Yes")
    new_entities = sum(
        1 for r in all_t1
        if r["already_in_database"] == "No" and r["independence_decision"] == "include"
    )
    ws.append(["Already in database tab", already_in])
    ws.append(["New independent entities (not in database)", new_entities])
    ws.append([])

    # Exclusion reason breakdown
    ws.append(["Exclusion reasons"])
    ws.cell(ws.max_row, 1).font = section_font
    reason_counts: dict[str, int] = {}
    for r in all_t1:
        if r["independence_decision"] == "exclude":
            reason = r["decision_reason"]
            reason_counts[reason] = reason_counts.get(reason, 0) + 1
    for reason, cnt in sorted(reason_counts.items(), key=lambda x: -x[1]):
        ws.append([f"  {reason}", cnt])
    ws.append([])

    # Match source breakdown
    ws.append(["CH number match sources"])
    ws.cell(ws.max_row, 1).font = section_font
    source_counts: dict[str, int] = {}
    for r in all_t1:
        src = r["match_source"] or "none"
        source_counts[src] = source_counts.get(src, 0) + 1
    for src, cnt in sorted(source_counts.items(), key=lambda x: -x[1]):
        ws.append([f"  {src}", cnt])
    ws.append([])

    # Task 2
    ws.append(["Task 2 — Reclassified excluded domiciliary rows"])
    ws.cell(ws.max_row, 1).font = section_font
    ws.append(["Total reclassified rows", len(task2_results)])
    t1_cnt = sum(1 for r in task2_results if r["new_tier"] == "Tier 1")
    t2_cnt = sum(1 for r in task2_results if r["new_tier"] == "Tier 2")
    t3_cnt = sum(1 for r in task2_results if r["new_tier"] == "Tier 3")
    nursing_cnt = sum(1 for r in task2_results if r["has_non_residential_nursing"] == "Yes")
    ws.append(["  → Tier 1 (director ≥ 60)", t1_cnt])
    ws.append(["  → Tier 2 (no director ≥ 60)", t2_cnt])
    ws.append(["  → Tier 3 (domiciliary + supported living)", t3_cnt])
    ws.append(["  Has non-residential nursing", nursing_cnt])

    ws.freeze_panes = "A3"


def write_output(
    task1_results: list[dict],
    task1_review: list[dict],
    task2_results: list[dict],
    out_path: Path,
):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Summary
    ws = wb.create_sheet("summary")
    build_summary(ws, task1_results, task1_review, task2_results)

    # Task 1 main
    t1_keys = [
        "provider_id", "provider_name", "postcode", "service_types",
        "local_authority", "region",
        "ownership_type", "brand_name", "charity_number",
        "candidate_company_number", "candidate_company_name", "company_type",
        "match_source", "independence_decision", "decision_reason",
        "already_in_database", "database_row_tier",
    ]
    t1_headers = [
        "Provider ID", "Provider Name", "Postcode", "Service Types",
        "Local Authority", "Region",
        "Ownership Type", "Brand Name", "Charity Number",
        "CH Company Number", "CH Company Name", "Company Type",
        "Match Source", "Decision", "Decision Reason",
        "Already in Database?", "Database Tier",
    ]
    ws = wb.create_sheet("task1_filtered_missing")
    write_table(ws, t1_headers, task1_results, t1_keys)
    style_decision_column(ws, "Decision")

    # Task 1 review queue
    ws = wb.create_sheet("task1_review_queue")
    write_table(ws, t1_headers, task1_review, t1_keys)
    style_decision_column(ws, "Decision")

    # Task 2
    t2_keys = [
        "company_number", "company_name", "postcode", "classification",
        "original_tier", "oldest_director", "directors_60_plus",
        "new_tier", "has_non_residential_nursing",
        "tier_reason", "data_quality_note",
    ]
    t2_headers = [
        "Company Number", "Company Name", "Postcode", "Classification",
        "Original Tier", "Oldest Director", "# Directors 60+",
        "New Tier", "Has Non-Residential Nursing?",
        "Tier Reason", "Data Quality Note",
    ]
    ws = wb.create_sheet("task2_reclassified")
    write_table(ws, t2_headers, task2_results, t2_keys)
    style_tier_column(ws, "New Tier")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# --------------- main ---------------


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default=str(INPUT_XLSX), help="Input workbook path")
    parser.add_argument("--db", default=str(DB_PATH), help="DuckDB path")
    parser.add_argument("--out", default=None, help="Output xlsx path")
    args = parser.parse_args()

    out_path = (
        Path(args.out) if args.out
        else Path("data/exports") / f"homecare_gap_analysis_{date.today():%Y%m%d}.xlsx"
    )

    print(f"Loading workbook: {args.input}")
    wb = openpyxl.load_workbook(args.input, read_only=True, data_only=True)

    missing_rows = load_potentially_missing(wb)
    print(f"  Potentially missing homecare: {len(missing_rows)} populated rows")

    db_rows = load_database_tab(wb)
    print(f"  Database tab: {len(db_rows)} rows")

    print(f"Connecting to DB: {args.db}")
    con = duckdb.connect(args.db, read_only=True)

    print("Running Task 1 — filter & match missing providers...")
    task1_results, task1_review = run_task1(missing_rows, db_rows, con)
    print(f"  Results: {len(task1_results)} decided, {len(task1_review)} in review queue")

    print("Running Task 2 — reclassify excluded domiciliary rows...")
    task2_results = run_task2(db_rows, con)
    print(f"  Reclassified: {len(task2_results)} rows")

    con.close()

    print(f"Writing output: {out_path}")
    write_output(task1_results, task1_review, task2_results, out_path)

    # Print summary
    all_t1 = task1_results + task1_review
    include = sum(1 for r in all_t1 if r["independence_decision"] == "include")
    exclude = sum(1 for r in all_t1 if r["independence_decision"] == "exclude")
    review = sum(1 for r in all_t1 if r["independence_decision"] == "review")
    t1_cnt = sum(1 for r in task2_results if r["new_tier"] == "Tier 1")
    t2_cnt = sum(1 for r in task2_results if r["new_tier"] == "Tier 2")
    t3_cnt = sum(1 for r in task2_results if r["new_tier"] == "Tier 3")

    print(f"\n=== Task 1 Summary ===")
    print(f"  Include: {include}  |  Exclude: {exclude}  |  Review: {review}")
    print(f"  Already in database: {sum(1 for r in all_t1 if r['already_in_database'] == 'Yes')}")
    print(f"\n=== Task 2 Summary ===")
    print(f"  Tier 1: {t1_cnt}  |  Tier 2: {t2_cnt}  |  Tier 3: {t3_cnt}")
    print(f"\nDone. Output: {out_path}")


if __name__ == "__main__":
    main()
